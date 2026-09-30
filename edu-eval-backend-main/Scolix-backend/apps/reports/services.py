"""
Génération de rapports.
- CSV : natif Python, toujours dispo
- Excel : openpyxl
- PDF : reportlab (à ajouter dans requirements.txt)

Le service retourne un HttpResponse prêt à streamer.
"""
import csv
import io
from django.http import HttpResponse
from django.utils import timezone

from apps.evaluations.models import EvaluationSubmission
from apps.sync.models import TeacherSync
from apps.attendance.models import AttendanceRecord
from apps.analytics.services import AnalyticsService


class ReportService:

    # ── CSV ──────────────────────────────────────────────────────────────────

    @staticmethod
    def teacher_ranking_csv(semester_id=None) -> HttpResponse:
        ranking = AnalyticsService.teacher_ranking(semester_id, top_n=200)
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="classement_enseignants.csv"'
        response.write("\ufeff")  # BOM pour Excel

        writer = csv.writer(response)
        writer.writerow(["Rang", "Enseignant", "Département", "Score moyen", "Nb évaluations"])
        for row in ranking:
            writer.writerow([row["rank"], row["teacher_name"], row["department"],
                              row["avg_score"], row["eval_count"]])
        return response

    @staticmethod
    def campaign_results_csv(campaign_id) -> HttpResponse:
        qs = EvaluationSubmission.objects.filter(
            campaign_id=campaign_id,
            status=EvaluationSubmission.Status.SUBMITTED,
        ).select_related("course", "course__department", "teacher", "student")

        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="resultats_campagne.csv"'
        response.write("\ufeff")

        writer = csv.writer(response)
        writer.writerow(["Cours", "Enseignant", "Département", "Score global", "Date soumission"])
        for s in qs:
            writer.writerow([
                s.course.code, s.teacher.full_name,
                s.course.department.name,
                float(s.global_score), s.submitted_at.strftime("%Y-%m-%d %H:%M"),
            ])
        return response

    @staticmethod
    def attendance_csv(teacher_id=None, period=None) -> HttpResponse:
        qs = AttendanceRecord.objects.select_related("teacher", "course")
        if teacher_id:
            qs = qs.filter(teacher_id=teacher_id)
        if period:
            year, month = period.split("-")
            qs = qs.filter(scheduled_at__year=int(year), scheduled_at__month=int(month))

        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="presences.csv"'
        response.write("\ufeff")

        writer = csv.writer(response)
        writer.writerow(["Enseignant", "Matricule", "Cours", "Date prévue", "Date réelle", "Statut", "Retard (min)", "Justifié"])
        for r in qs.order_by("-scheduled_at"):
            writer.writerow([
                r.teacher.full_name, r.teacher.matricule, r.course.code,
                r.scheduled_at.strftime("%Y-%m-%d %H:%M"),
                r.actual_at.strftime("%Y-%m-%d %H:%M") if r.actual_at else "",
                r.status, r.delay_minutes, "Oui" if r.is_justified else "Non",
            ])
        return response

    # ── Excel (openpyxl) ─────────────────────────────────────────────────────

    @staticmethod
    def teacher_ranking_excel(semester_id=None) -> HttpResponse:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise ImportError("openpyxl est requis. Lancez : pip install openpyxl")

        ranking = AnalyticsService.teacher_ranking(semester_id, top_n=200)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Classement Enseignants"

        headers = ["Rang", "Enseignant", "Département", "Score moyen", "Nb évaluations"]
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_idx, r in enumerate(ranking, 2):
            ws.append([r["rank"], r["teacher_name"], r["department"], r["avg_score"], r["eval_count"]])
            if r["avg_score"] >= 75:
                ws.cell(row=row_idx, column=4).fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            elif r["avg_score"] < 55:
                ws.cell(row=row_idx, column=4).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 20

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="classement_enseignants.xlsx"'
        return response

    # ── PDF (reportlab) ───────────────────────────────────────────────────────

    @staticmethod
    def teacher_summary_pdf(teacher_id, semester_id=None) -> HttpResponse:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
        except ImportError:
            raise ImportError("reportlab est requis. Lancez : pip install reportlab")

        from apps.analytics.services import AnalyticsService

        try:
            teacher = TeacherSync.objects.get(id=teacher_id)
        except TeacherSync.DoesNotExist:
            raise ValueError("Enseignant introuvable.")

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph(f"Rapport — {teacher.full_name}", styles["Title"]))
        elements.append(Paragraph(f"Département : {teacher.department.name}", styles["Normal"]))
        elements.append(Paragraph(f"Matricule : {teacher.matricule}", styles["Normal"]))
        elements.append(Spacer(1, 20))

        ranking = AnalyticsService.teacher_ranking(semester_id, top_n=200)
        teacher_rank = next((r for r in ranking if r["teacher_id"] == str(teacher_id)), None)

        if teacher_rank:
            elements.append(Paragraph(f"Score moyen : {teacher_rank['avg_score']} / 100", styles["Heading2"]))
            elements.append(Paragraph(f"Rang : {teacher_rank['rank']}", styles["Normal"]))
            elements.append(Paragraph(f"Nombre d'évaluations : {teacher_rank['eval_count']}", styles["Normal"]))

        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="rapport_{teacher.matricule}.pdf"'
        return response
